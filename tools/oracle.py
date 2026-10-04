#!/usr/bin/env python3
"""Independent, fixture-specific XLSX integrity and formula-behavior oracle.

No application modules, JavaScript evaluation, or cached spreadsheet values are
used by the evaluator. This is intentionally not a general Excel evaluator.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/package/2006/relationships",
      "ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
OUTPUT_CELLS = ("C2", "C3", "C4")
STATES = tuple((x, y) for x in (0, 1, 8, 10) for y in (0, 2, 3))
SOURCE_FORMULAS = {
    ("Plan", "C2"): "=B2*Derived!$B$2+Constants!$B$3",
    ("Plan", "C3"): "=B3*Constants!$B$4",
    ("Plan", "C4"): "=SUM(C2:C3)",
    ("Derived", "B2"): "=Constants!B2*2",
}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def expected_values(x: int, y: int) -> dict[str, int]:
    """Closed-form expectations authored separately from the app/evaluator."""
    return {"C2": 6 * x - 2, "C3": 5 * y, "C4": 6 * x + 5 * y - 2}


def check_values(actual: dict, expected: dict, context: str) -> None:
    require(set(actual) == set(expected), f"{context}: wrong result cells")
    for address, value in actual.items():
        require(isinstance(value, (float, int)) and not isinstance(value, bool)
                and math.isfinite(value), f"{context} {address}: non-numeric result {value!r}")
        require(math.isclose(value, expected[address], rel_tol=0, abs_tol=1e-10),
                f"{context} {address}: {value!r} != {expected[address]!r}")


# Reference-first tokenization allows Constants!B2 without treating Constants as
# a name. Quoted sheets are accepted; named ranges and arbitrary functions are not.
TOKEN = re.compile(
    r"(?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!)?\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)(?![A-Za-z0-9_])"
    r"|(?P<number>(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?)"
    r"|(?P<name>[A-Za-z_][A-Za-z_0-9.]*)|(?P<op>[+*/(),:\-])"
)


class FormulaParser:
    """Small recursive-descent grammar: numbers, A1 refs, + - * /, SUM."""

    def __init__(self, formula: str):
        require(formula.startswith("="), "Formula must begin with =")
        self.tokens = []
        position = 1
        while position < len(formula):
            if formula[position].isspace():
                position += 1
                continue
            token = TOKEN.match(formula, position)
            require(token is not None, f"Unsupported formula token at {formula[position:]!r}")
            self.tokens.append((token.lastgroup, token.group()))
            position = token.end()
        self.tokens.append(("end", ""))
        self.index = 0

    def peek(self) -> tuple[str, str]:
        return self.tokens[self.index]

    def take(self, value: str | None = None) -> tuple[str, str]:
        token = self.peek()
        require(value is None or token[1] == value, f"Expected {value!r}, got {token!r}")
        self.index += 1
        return token

    def parse(self):
        result = self.expression()
        require(self.peek()[0] == "end", f"Unconsumed formula token {self.peek()!r}")
        return result

    def expression(self):
        node = self.product()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            node = (op, node, self.product())
        return node

    def product(self):
        node = self.unary()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            node = (op, node, self.unary())
        return node

    def unary(self):
        if self.peek()[1] in ("+", "-"):
            return ("unary", self.take()[1], self.unary())
        return self.primary()

    def primary(self):
        kind, text = self.take()
        if kind == "number":
            value = float(text)
            require(math.isfinite(value), "Non-finite numeric literal")
            return ("number", value)
        if kind == "ref":
            first = ("ref", text)
            if self.peek()[1] == ":":
                self.take(":")
                last = self.take()
                require(last[0] == "ref", "Range must end in a cell reference")
                return ("range", text, last[1])
            return first
        if text == "(":
            node = self.expression()
            self.take(")")
            return node
        if kind == "name":
            require(text.upper() == "SUM", f"Unsupported function {text!r}")
            self.take("(")
            arguments = [self.expression()]
            while self.peek()[1] == ",":
                self.take(",")
                arguments.append(self.expression())
            self.take(")")
            return ("sum", arguments)
        raise AssertionError(f"Unexpected formula token {(kind, text)!r}")


def split_reference(text: str, current_sheet: str) -> tuple[str, str]:
    if "!" in text:
        sheet, address = text.rsplit("!", 1)
        if sheet.startswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
    else:
        sheet, address = current_sheet, text
    return sheet, address.replace("$", "").upper()


def expand_range(first: str, last: str, current_sheet: str) -> list[tuple[str, str]]:
    sheet1, address1 = split_reference(first, current_sheet)
    sheet2, address2 = split_reference(last, sheet1)
    require(sheet1 == sheet2, "Cross-sheet ranges are outside the oracle grammar")
    col1, row1 = coordinate_from_string(address1)
    col2, row2 = coordinate_from_string(address2)
    left, right = column_index_from_string(col1), column_index_from_string(col2)
    require(left <= right and row1 <= row2, "Descending range is not supported")
    require((right - left + 1) * (row2 - row1 + 1) <= 1000, "Oracle range limit exceeded")
    return [(sheet1, f"{get_column_letter(col)}{row}")
            for row in range(row1, row2 + 1) for col in range(left, right + 1)]


def references(node, sheet: str) -> set[tuple[str, str]]:
    kind = node[0]
    if kind == "ref":
        return {split_reference(node[1], sheet)}
    if kind == "range":
        return set(expand_range(node[1], node[2], sheet))
    if kind == "number":
        return set()
    if kind == "unary":
        return references(node[2], sheet)
    if kind == "sum":
        return set().union(*(references(argument, sheet) for argument in node[1]))
    return references(node[1], sheet) | references(node[2], sheet)


class FormulaEvaluator:
    def __init__(self, workbook, overrides: dict | None = None):
        self.workbook = workbook
        self.overrides = overrides or {}
        self.active = set()
        self.memo = {}

    @staticmethod
    def scalar(value):
        require(not isinstance(value, list), "Range used outside SUM")
        require(isinstance(value, (float, int)) and not isinstance(value, bool)
                and math.isfinite(value), f"Non-numeric oracle operand {value!r}")
        return value

    def cell(self, sheet: str, address: str):
        key = (sheet, address)
        if key in self.overrides:
            return self.scalar(self.overrides[key])
        if key in self.memo:
            return self.memo[key]
        require(key not in self.active, f"Circular oracle reference {key}")
        require(sheet in self.workbook.sheetnames, f"Missing sheet {sheet!r}")
        cell = self.workbook[sheet][address]
        self.active.add(key)
        try:
            result = self.node(FormulaParser(cell.value).parse(), sheet) if cell.data_type == "f" else cell.value
            self.memo[key] = self.scalar(result)
            return self.memo[key]
        finally:
            self.active.remove(key)

    def node(self, node, sheet: str):
        kind = node[0]
        if kind == "number":
            return node[1]
        if kind == "ref":
            return self.cell(*split_reference(node[1], sheet))
        if kind == "range":
            return [self.cell(*key) for key in expand_range(node[1], node[2], sheet)]
        if kind == "unary":
            value = self.scalar(self.node(node[2], sheet))
            return value if node[1] == "+" else -value
        if kind == "sum":
            values = [self.node(argument, sheet) for argument in node[1]]
            return sum(self.scalar(v) for item in values for v in (item if isinstance(item, list) else [item]))
        left, right = self.scalar(self.node(node[1], sheet)), self.scalar(self.node(node[2], sheet))
        if kind == "+":
            return left + right
        if kind == "-":
            return left - right
        if kind == "*":
            return left * right
        require(kind == "/" and right != 0, "Invalid operation or division by zero")
        return left / right


def evaluate(workbook, x: int, y: int) -> dict:
    engine = FormulaEvaluator(workbook, {("Plan", "B2"): x, ("Plan", "B3"): y})
    return {address: engine.cell("Plan", address) for address in OUTPUT_CELLS}


def relationship_target(part: str, target: str) -> str:
    target = unquote(target)
    require("\\" not in target and not re.match(r"^[a-zA-Z]+:", target), "Invalid internal relationship target")
    if target.startswith("/"):
        normalized = posixpath.normpath(target.lstrip("/"))
    else:
        base = "" if part == "_rels/.rels" else posixpath.dirname(posixpath.dirname(part))
        normalized = posixpath.normpath(posixpath.join(base, target))
    require(not normalized.startswith("../"), "Relationship escapes package root")
    return normalized


def inspect_package(path: Path, *, detached: bool) -> dict:
    with zipfile.ZipFile(path) as archive:
        names = [entry.filename for entry in archive.infolist() if not entry.is_dir()]
        require(len(names) == len(set(names)), "Duplicate ZIP entries")
        require(archive.testzip() is None, "ZIP CRC failure")
        require(all(not name.startswith("/") and ".." not in name.split("/") for name in names), "Unsafe package path")
        parts = {name: archive.read(name) for name in names}
    for required in ("[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml"):
        require(required in parts, f"Missing package part {required}")
    xml = {name: ET.fromstring(data) for name, data in parts.items() if name.endswith((".xml", ".rels"))}
    relationships = []
    for name, root in xml.items():
        if not name.endswith(".rels"):
            continue
        ids = [rel.get("Id") for rel in root]
        require(len(ids) == len(set(ids)), f"Duplicate relationship IDs in {name}")
        for rel in root:
            require(rel.get("TargetMode") != "External", f"Unexpected external relationship in {name}")
            target = relationship_target(name, rel.get("Target", ""))
            require(target in parts, f"Dangling relationship {name} -> {target}")
            relationships.append({"part": name, "id": rel.get("Id"), "type": rel.get("Type"), "target": target})
    workbook_rels = {rel["id"]: rel for rel in relationships if rel["part"] == "xl/_rels/workbook.xml.rels"}
    sheets = {}
    for sheet in xml["xl/workbook.xml"].findall("s:sheets/s:sheet", NS):
        rel = workbook_rels[sheet.get(f"{{{NS['r']}}}id")]
        require(rel["type"].endswith("/worksheet"), "Workbook sheet points to a non-worksheet")
        require(sheet.get("name") not in sheets, "Duplicate sheet name")
        sheets[sheet.get("name")] = rel["target"]
    expected_names = {"Plan"} if detached else {"Plan", "Constants", "Derived"}
    require(set(sheets) == expected_names, f"Unexpected sheets: {list(sheets)}")
    actual_worksheet_parts = {name for name in parts if re.fullmatch(r"xl/worksheets/[^/]+\.xml", name)}
    require(actual_worksheet_parts == set(sheets.values()), "Orphan or missing worksheet payload")
    worksheet_rels = {rel["target"] for rel in workbook_rels.values() if rel["type"].endswith("/worksheet")}
    require(worksheet_rels == set(sheets.values()), "Orphan worksheet relationship")
    overrides = xml["[Content_Types].xml"].findall("ct:Override", NS)
    require(all(item.get("PartName", "").lstrip("/") in parts for item in overrides), "Dangling content-type override")
    manifest_sheets = {item.get("PartName").lstrip("/") for item in overrides if item.get("ContentType", "").endswith("worksheet+xml")}
    require(manifest_sheets == set(sheets.values()), "Worksheet manifest mismatch")
    formula_records = {}
    for sheet, part in sheets.items():
        for cell in xml[part].findall(".//s:sheetData/s:row/s:c", NS):
            formula = cell.find("s:f", NS)
            if formula is not None:
                require(not formula.attrib, f"Unsupported formula attributes at {sheet}!{cell.get('r')}")
                cache = cell.find("s:v", NS)
                formula_records[f"{sheet}!{cell.get('r')}"] = {
                    "formula": "=" + (formula.text or ""),
                    "cache": None if cache is None or cache.text is None else cache.text,
                    "type": cell.get("t", "n"),
                    "style": cell.get("s", "0"),
                }
    expected_formulas = {f"Plan!{address}" for address in OUTPUT_CELLS} if detached else {f"{s}!{a}" for s, a in SOURCE_FORMULAS}
    require(set(formula_records) == expected_formulas, f"Formula cells missing or unexpected: {list(formula_records)}")
    if detached:
        for name, data in parts.items():
            require(b"Constants" not in data and b"Derived" not in data, f"Removed sheet name leaked into {name}")
            require(not any(marker in name.lower() for marker in ("externallink", "calcchain", "vba", "customxml")), f"Unexpected retained package part {name}")
        initial = expected_values(8, 2)
        for key, record in formula_records.items():
            require(record["type"] not in ("e", "str", "inlineStr", "s"), f"Invalid formula result type {key}")
            if record["cache"] is not None:
                require(float(record["cache"]) == initial[key.split("!")[1]], f"Stale formula cache at {key}")
        calc = xml["xl/workbook.xml"].find("s:calcPr", NS)
        require(calc is not None and calc.get("calcMode", "auto") == "auto", "Automatic recalculation not enabled")
        require(calc.get("fullCalcOnLoad") in ("1", "true") or calc.get("forceFullCalc") in ("1", "true"), "Full recalculation not requested")
    else:
        for key, record in formula_records.items():
            require(record["cache"] is not None and float(record["cache"]) == 999, f"Source cache must be deliberately stale (999) at {key}")
    return {"sheets": sheets, "parts": sorted(parts), "relationships": relationships,
            "formula_records": formula_records, "styles_sha256": hashlib.sha256(parts["xl/styles.xml"]).hexdigest()}


def validate_artifacts(source: Path, handoff: Path) -> dict:
    require(source.is_file(), f"Missing source fixture: {source}")
    require(handoff.is_file(), f"Missing exported workbook: {handoff}")
    source_package = inspect_package(source, detached=False)
    handoff_package = inspect_package(handoff, detached=True)
    original = openpyxl.load_workbook(source, data_only=False)
    output = openpyxl.load_workbook(handoff, data_only=False)
    for (sheet, address), formula in SOURCE_FORMULAS.items():
        require(original[sheet][address].value == formula, f"Unexpected source formula {sheet}!{address}")
    for sheet, address, value in (("Plan", "B2", 8), ("Plan", "B3", 2), ("Constants", "B2", 3),
                                  ("Constants", "B3", -2), ("Constants", "B4", 5)):
        require(original[sheet][address].value == value, f"Unexpected source input {sheet}!{address}")
    original_cells = {cell.coordinate: cell for row in original["Plan"] for cell in row if cell.value is not None or cell.has_style}
    output_cells = {cell.coordinate: cell for row in output["Plan"] for cell in row if cell.value is not None or cell.has_style}
    require(set(original_cells) == set(output_cells), "Kept-sheet cell set changed or extra constants were inserted")
    require(source_package["styles_sha256"] == handoff_package["styles_sha256"], "styles.xml changed")
    styled_cells = []
    for address, before in original_cells.items():
        after = output_cells[address]
        require(before._style == after._style, f"Cell style changed at Plan!{address}")
        require(before.number_format == after.number_format, f"Number format changed at Plan!{address}")
        if before.has_style:
            styled_cells.append(address)
        if address not in OUTPUT_CELLS:
            require(before.value == after.value and before.data_type == after.data_type, f"Non-formula cell changed at Plan!{address}")
    require(styled_cells, "Fixture needs at least one non-default style to exercise preservation")
    def preserved_sheet_structure(path, part):
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read(part))
        for cell in root.findall(".//s:sheetData/s:row/s:c", NS):
            if cell.get("r") in OUTPUT_CELLS:
                for child in list(cell):
                    if child.tag in (f"{{{NS['s']}}}f", f"{{{NS['s']}}}v"):
                        cell.remove(child)
        def canonical(element):
            return (element.tag, sorted(element.attrib.items()),
                    element.text if element.text and element.text.strip() else None,
                    [canonical(child) for child in element])
        return canonical(root)
    require(preserved_sheet_structure(source, source_package["sheets"]["Plan"])
            == preserved_sheet_structure(handoff, handoff_package["sheets"]["Plan"]),
            "Kept-sheet structure, dimensions, styles, panes, merges, or content changed")
    formulas = {address: output["Plan"][address].value for address in OUTPUT_CELLS}
    dependencies = {}
    expected_dependencies = {"C2": {("Plan", "B2")}, "C3": {("Plan", "B3")}, "C4": {("Plan", "C2"), ("Plan", "C3")}}
    for address, formula in formulas.items():
        require(output["Plan"][address].data_type == "f", f"Formula replaced with value at Plan!{address}")
        deps = references(FormulaParser(formula).parse(), "Plan")
        require(deps == expected_dependencies[address], f"Editable dependency changed at {address}: {deps}")
        dependencies[address] = sorted(f"{sheet}!{cell}" for sheet, cell in deps)
    require(re.sub(r"\s+", "", formulas["C4"]).upper() == "=SUM(C2:C3)", "Kept-only SUM formula changed")
    states = []
    for x, y in STATES:
        expected = expected_values(x, y)
        source_values = evaluate(original, x, y)
        handoff_values = evaluate(output, x, y)
        check_values(source_values, expected, f"source x={x}, y={y}")
        check_values(handoff_values, expected, f"handoff x={x}, y={y}")
        states.append({"inputs": {"B2": x, "B3": y}, "expected": expected,
                       "source_formula_values": source_values, "handoff_formula_values": handoff_values})
    return {"status": "passed", "method": "independent Python formula parser + openpyxl + ZIP/XML inspection",
            "native_application_recalculation": False, "source": {"path": str(source), "sha256": sha256(source)},
            "handoff": {"path": str(handoff), "sha256": sha256(handoff)},
            "source_package": source_package, "handoff_package": handoff_package,
            "formulas": formulas, "local_dependencies": dependencies, "preserved_styled_cells": styled_cells,
            "initial_values": evaluate(output, 8, 2), "mutated_values": evaluate(output, 10, 3),
            "closed_form_states": states, "state_count": len(states)}


def artifact_mutation_tests(source: Path, handoff: Path) -> list[dict]:
    """Prove the fixture checks reject plausible bad exports, not just pass one."""
    package = inspect_package(handoff, detached=True)
    sheet_part = package["sheets"]["Plan"]
    with zipfile.ZipFile(handoff) as archive:
        original_parts = {name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}
    cases = ("flatten_formula", "wrong_operator", "stale_cache", "removed_name_leak",
             "orphan_sheet", "dangling_relationship", "changed_cell_style", "missing_manifest_sheet")
    results = []
    with tempfile.TemporaryDirectory(prefix="sheet-detach-oracle-") as temporary:
        for case in cases:
            parts = dict(original_parts)
            root = ET.fromstring(parts[sheet_part])
            cell = next(c for c in root.findall(".//s:c", NS) if c.get("r") == "C2")
            if case == "flatten_formula":
                cell.remove(cell.find("s:f", NS))
            elif case == "wrong_operator":
                cell.find("s:f", NS).text = "B2+(6)+(-2)"
            elif case == "stale_cache":
                value = cell.find("s:v", NS)
                if value is None:
                    value = ET.SubElement(cell, f"{{{NS['s']}}}v")
                value.text = "999"
            elif case == "changed_cell_style":
                cell.set("s", "0")
            elif case == "removed_name_leak":
                parts["xl/workbook.xml"] += b"<!-- Constants -->"
            elif case == "orphan_sheet":
                parts["xl/worksheets/forgotten.xml"] = parts[sheet_part]
            elif case == "dangling_relationship":
                rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
                next(rel for rel in rels if rel.get("Type", "").endswith("/worksheet")).set("Target", "worksheets/missing.xml")
                parts["xl/_rels/workbook.xml.rels"] = ET.tostring(rels)
            elif case == "missing_manifest_sheet":
                manifest = ET.fromstring(parts["[Content_Types].xml"])
                for item in list(manifest):
                    if item.get("ContentType", "").endswith("worksheet+xml"):
                        manifest.remove(item)
                parts["[Content_Types].xml"] = ET.tostring(manifest)
            if case in ("flatten_formula", "wrong_operator", "stale_cache", "changed_cell_style"):
                parts[sheet_part] = ET.tostring(root)
            path = Path(temporary) / f"{case}.xlsx"
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            try:
                validate_artifacts(source, path)
            except AssertionError as exc:
                results.append({"mutation": case, "status": "rejected", "reason": str(exc)})
            else:
                raise AssertionError(f"Oracle failed to reject mutation {case}")
    return results


def parser_self_test() -> dict:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Plan"
    sheet["B2"], sheet["B3"] = 8, 2
    checks = {
        "=B2*(6)+(-2)": 46, "=B3*(5)": 10, "=2+3*4": 14,
        "=(2+3)*4": 20, "=10-3-2": 5, "=12/3/2": 2,
        "=--2": 2, "=SUM(B2:B3,(-2),3*4)": 20,
        "=SUM('Plan'!$B$2:$B$3)": 10, "=1e2+0.5": 100.5,
    }
    for formula, expected in checks.items():
        actual = FormulaEvaluator(book).node(FormulaParser(formula).parse(), "Plan")
        require(actual == expected, f"Parser self-test failed: {formula}")
    rejected = ("=INDIRECT(\"B2\")", "=B2garbage", "=2^3", "=2+", "=SUM(B2:B3)junk", "=[other.xlsx]Plan!B2", "=1/0")
    for formula in rejected:
        try:
            FormulaEvaluator(book).node(FormulaParser(formula).parse(), "Plan")
        except (AssertionError, ValueError):
            continue
        raise AssertionError(f"Oracle accepted unsupported/invalid formula {formula}")
    return {"valid_formulas": len(checks), "rejected_formulas": len(rejected)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "generated/source.xlsx")
    parser.add_argument("--input", type=Path, default=ROOT / "generated/handoff.xlsx")
    parser.add_argument("--out", type=Path, default=ROOT / "test-results/oracle/evidence.json")
    parser.add_argument("--self-test", action="store_true", help="Exercise independent parser without input artifacts")
    args = parser.parse_args()
    try:
        self_tests = parser_self_test()
        if args.self_test:
            print(json.dumps({"status": "passed", "parser_self_tests": self_tests}))
            return 0
        result = validate_artifacts(args.source, args.input)
        result["parser_self_tests"] = self_tests
        result["rejected_artifact_mutations"] = artifact_mutation_tests(args.source, args.input)
    except Exception as exc:
        result = {"status": "failed", "error": f"{type(exc).__name__}: {exc}", "native_application_recalculation": False}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"status": result["status"], "evidence": str(args.out), "state_count": result.get("state_count", 0), "error": result.get("error")}))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    sys.exit(main())
