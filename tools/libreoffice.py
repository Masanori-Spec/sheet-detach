#!/usr/bin/env python3
"""Hosted-runner actual-consumer test of the browser-downloaded XLSX.

Requires a GitHub-hosted Linux runner with LibreOffice Calc. --validate performs
local syntax/configuration checks only and NEVER launches LibreOffice.
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.workbook.properties import CalcProperties

from oracle import (OUTPUT_CELLS, ROOT, STATES, check_values, expected_values,
                    parser_self_test, require, sha256, validate_artifacts)

BROWSER_DOWNLOAD = ROOT / "test-results/browser/actual-artifacts/handoff.xlsx"
# LibreOffice's documented OOXML load setting: 0 = always recalculate.
# This is isolated test-application configuration, never the user's profile.
PROFILE_XML = '''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
    <prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
</oor:items>
'''


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def conversion_command(binary: str, profile: Path, source: Path, destination: Path) -> list[str]:
    return [binary, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless",
            "--nologo", "--nodefault", "--nofirststartwizard", "--norestore",
            "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(destination), str(source)]


def run_consumer(binary: str, source: Path, destination: Path, scratch: Path, label: str, runs: list) -> Path:
    require(source.is_file(), f"Missing native consumer input {source}")
    destination.mkdir(parents=True, exist_ok=True)
    output = destination / source.name
    require(source.resolve() != output.resolve(), "Consumer must save to a separate directory")
    require(not output.exists(), f"Refusing to reuse stale native output {output}")
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-", dir=scratch))
    (profile / "user").mkdir()
    (profile / "user/registrymodifications.xcu").write_text(PROFILE_XML)
    command = conversion_command(binary, profile, source, destination)
    started = time.monotonic()
    record = {"label": label, "source_sha256": sha256(source), "fresh_profile": True,
              "recalculate_on_load": "always", "status": "running"}
    runs.append(record)
    try:
        environment = dict(os.environ, SAL_USE_VCLPLUGIN="gen", LANG="C.UTF-8", LC_ALL="C.UTF-8")
        result = subprocess.run(command, capture_output=True, text=True, timeout=90, env=environment, check=False)
        record.update(returncode=result.returncode, stdout=result.stdout, stderr=result.stderr)
        require(result.returncode == 0, f"LibreOffice failed for {label}: {result.returncode}: {result.stderr}")
        require(output.is_file() and output.stat().st_size > 0, f"LibreOffice produced no saved XLSX for {label}: {result.stdout} {result.stderr}")
        record.update(status="saved", output_sha256=sha256(output), output=str(output))
        return output
    except Exception as exc:
        record.update(status="failed", error=f"{type(exc).__name__}: {exc}")
        raise
    finally:
        record["duration_seconds"] = round(time.monotonic() - started, 3)
        shutil.rmtree(profile, ignore_errors=True)


def read_native_values(path: Path, x: int, y: int) -> dict:
    """Only read LibreOffice's saved results; openpyxl does not calculate them."""
    values_book = openpyxl.load_workbook(path, data_only=True)
    formulas_book = openpyxl.load_workbook(path, data_only=False)
    require(values_book.sheetnames == ["Plan"] and formulas_book.sheetnames == ["Plan"], "Native round-trip changed the kept sheets")
    sheet = formulas_book["Plan"]
    require(sheet["B2"].value == x and sheet["B3"].value == y, "Native round-trip changed authored input values")
    formulas = {}
    for address in OUTPUT_CELLS:
        cell = sheet[address]
        require(cell.data_type == "f", f"Native consumer flattened formula {address}")
        require(not any(marker in cell.value for marker in ("#REF!", "Constants", "Derived", "[")), f"Native formula has invalid dependencies at {address}")
        formulas[address] = cell.value
    actual = {address: values_book["Plan"][address].value for address in OUTPUT_CELLS}
    check_values(actual, expected_values(x, y), f"LibreOffice {path}")
    return {"path": str(path), "sha256": sha256(path), "values": actual, "formulas": formulas}


def author_mutation(base: Path, destination: Path, x: int, y: int) -> dict:
    """Author a disposable consumer-test copy; never evaluate or fake caches."""
    book = openpyxl.load_workbook(base, data_only=False)
    require(book.sheetnames == ["Plan"], "Only the detached handoff may be mutated")
    sheet = book["Plan"]
    before_formulas = {address: sheet[address].value for address in OUTPUT_CELLS}
    sheet["B2"] = x
    sheet["B3"] = y
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    reopened = openpyxl.load_workbook(destination, data_only=False)
    require({address: reopened["Plan"][address].value for address in OUTPUT_CELLS} == before_formulas,
            "Authoring test input changed formulas")
    empty_cache = openpyxl.load_workbook(destination, data_only=True)
    require(all(empty_cache["Plan"][address].value is None for address in OUTPUT_CELLS),
            "Mutation authoring must clear caches, not manufacture expected results")
    return {"author": "openpyxl (test copies only; no calculation)", "path": str(destination),
            "sha256": sha256(destination), "inputs": {"B2": x, "B3": y}, "formula_caches_before_consumer": None}


def validate_configuration() -> dict:
    for filename in (Path(__file__), ROOT / "tools/oracle.py"):
        compile(filename.read_text(), str(filename), "exec")
    ET.fromstring(PROFILE_XML)
    require(len(STATES) == 12 and len(set(STATES)) == 12, "State coverage changed")
    require((8, 2) in STATES and (10, 3) in STATES, "Initial and changed states must be included")
    command = conversion_command("libreoffice", Path("/tmp/profile"), Path("/tmp/source.xlsx"), Path("/tmp/output"))
    require("--headless" in command and "xlsx:Calc MS Excel 2007 XML" in command, "Invalid conversion command")
    return {"status": "configuration_validated", "native_application_recalculation": False,
            "libreoffice_launched": False, "state_count": len(STATES), "parser_self_tests": parser_self_test()}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "generated/source.xlsx")
    parser.add_argument("--input", type=Path, default=BROWSER_DOWNLOAD)
    parser.add_argument("--out", type=Path, default=ROOT / "test-results/native/evidence.json")
    parser.add_argument("--validate", action="store_true", help="Local syntax/config validation only; never launches LibreOffice")
    args = parser.parse_args()
    if args.validate:
        print(json.dumps(validate_configuration()))
        return 0
    report = {"status": "running", "started_at": utc_now(), "consumer": "LibreOffice Calc",
              "native_application_recalculation": False, "runs": [], "closed_form_states": [],
              "environment": {"platform": platform.platform(), "python": platform.python_version(),
                              "openpyxl": openpyxl.__version__, "github_sha": os.environ.get("GITHUB_SHA"),
                              "github_run_id": os.environ.get("GITHUB_RUN_ID"),
                              "runner_environment": os.environ.get("RUNNER_ENVIRONMENT")}}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    try:
        require(os.environ.get("GITHUB_ACTIONS") == "true"
                and os.environ.get("RUNNER_ENVIRONMENT") == "github-hosted" and sys.platform == "linux",
                "Native launch is restricted to GitHub-hosted Linux runners. Use --validate for local checks.")
        require(args.input.resolve() == BROWSER_DOWNLOAD.resolve(),
                "Native input must be the actual browser download at test-results/browser/actual-artifacts/handoff.xlsx")
        require(args.input.is_file(), "Actual browser download is missing; generated fixture fallback is forbidden")
        binary = shutil.which("libreoffice") or shutil.which("soffice")
        require(binary is not None, "LibreOffice Calc is required; install official Ubuntu libreoffice-calc in the hosted workflow")
        validation = validate_artifacts(args.source, args.input)
        report["artifact_preflight"] = validation
        report["browser_download_sha256"] = sha256(args.input)
        version = subprocess.run([binary, "--version"], capture_output=True, text=True, timeout=15, check=True)
        report["consumer_version"] = version.stdout.strip() or version.stderr.strip()
        require("LibreOffice" in report["consumer_version"], "Could not establish native consumer version")
        # Unique evidence folder avoids confusing a previous run's XLSX with this run.
        run_root = Path(tempfile.mkdtemp(prefix="consumer-", dir=args.out.parent))
        report["artifacts_directory"] = str(run_root)
        with tempfile.TemporaryDirectory(prefix="sheet-detach-native-") as temporary:
            scratch = Path(temporary)
            initial_input = run_root / "initial/input/handoff.xlsx"
            initial_input.parent.mkdir(parents=True)
            shutil.copy2(args.input, initial_input)
            require(sha256(initial_input) == report["browser_download_sha256"], "Initial native input is not byte-identical to browser download")
            saved = run_consumer(binary, initial_input, run_root / "initial/saved", scratch, "initial save", report["runs"])
            first = read_native_values(saved, 8, 2)
            reopened = run_consumer(binary, saved, run_root / "initial/reopened", scratch, "initial reopen/save", report["runs"])
            second = read_native_values(reopened, 8, 2)
            report["initial"] = {"inputs": {"B2": 8, "B3": 2}, "expected": expected_values(8, 2), "saved": first, "reopened": second}
            for x, y in STATES:
                state_root = run_root / f"x{x}-y{y}"
                test_input = state_root / "input/handoff.xlsx"
                authored = author_mutation(reopened, test_input, x, y)
                saved = run_consumer(binary, test_input, state_root / "saved", scratch, f"x{x}-y{y} save", report["runs"])
                first = read_native_values(saved, x, y)
                reread = run_consumer(binary, saved, state_root / "reopened", scratch, f"x{x}-y{y} reopen/save", report["runs"])
                second = read_native_values(reread, x, y)
                report["closed_form_states"].append({"inputs": {"B2": x, "B3": y}, "expected": expected_values(x, y),
                                                       "authoring": authored, "saved": first, "reopened": second})
                # Keep useful partial evidence if the next native process fails.
                args.out.write_text(json.dumps(report, indent=2) + "\n")
        require(sha256(args.input) == report["browser_download_sha256"], "Browser download was modified during native test")
        require(len(report["closed_form_states"]) == 12 and len(report["runs"]) == 26, "Incomplete native state or reopen coverage")
        report.update(status="passed", native_application_recalculation=True, state_count=12,
                      native_conversion_count=26, original_browser_download_unchanged=True)
    except Exception as exc:
        report.update(status="failed", error=f"{type(exc).__name__}: {exc}")
    report["finished_at"] = utc_now()
    args.out.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps({"status": report["status"], "evidence": str(args.out), "consumer_version": report.get("consumer_version"),
                      "state_count": report.get("state_count", 0), "error": report.get("error")}))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    sys.exit(main())
